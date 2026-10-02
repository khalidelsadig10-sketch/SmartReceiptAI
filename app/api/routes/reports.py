from datetime import date
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy.orm import Session

from app.api.routes.auth import get_current_user
from app.core.database import SessionLocal
from app.models.user import UserModel
from app.schemas.reports import ReportResponse
from app.services.reports_service import ReportsService


router = APIRouter(
    prefix="/api/v1/reports",
    tags=["Reports"],
)


def get_db():
    db = SessionLocal()

    try:
        yield db
    finally:
        db.close()


def get_report_data(
    db: Session,
    current_user: UserModel,
    start_date: date,
    end_date: date,
):
    return ReportsService.get_report(
        db=db,
        user_id=current_user.id,
        start_date=start_date,
        end_date=end_date,
    )


@router.get(
    "/",
    response_model=ReportResponse,
)
def get_report(
    start_date: date,
    end_date: date,
    db: Session = Depends(get_db),
    current_user: UserModel = Depends(
        get_current_user
    ),
):
    try:
        report = get_report_data(
            db=db,
            current_user=current_user,
            start_date=start_date,
            end_date=end_date,
        )

        return {
            "success": True,
            "data": report,
        }

    except ValueError as exc:
        raise HTTPException(
            status_code=400,
            detail=str(exc),
        ) from exc

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail="Failed to generate report.",
        ) from exc


@router.get(
    "/export/excel",
)
def export_report_excel(
    start_date: date,
    end_date: date,
    db: Session = Depends(get_db),
    current_user: UserModel = Depends(
        get_current_user
    ),
):
    try:
        report = get_report_data(
            db=db,
            current_user=current_user,
            start_date=start_date,
            end_date=end_date,
        )

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Report"

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="D9E2F3",
        )

        header_font = Font(
            bold=True,
        )

        worksheet["A1"] = "SmartReceiptAI Report"
        worksheet["A1"].font = Font(
            bold=True,
            size=16,
        )

        worksheet["A2"] = "Start Date"
        worksheet["B2"] = report.start_date

        worksheet["A3"] = "End Date"
        worksheet["B3"] = report.end_date

        worksheet["A4"] = "Generated At"
        worksheet["B4"] = report.generated_at

        worksheet["A6"] = "Overview"
        worksheet["A6"].font = header_font
        worksheet["A6"].fill = header_fill

        overview_rows = [
            [
                "Total Receipts",
                report.overview.total_receipts,
            ],
            [
                "Total Items",
                report.overview.total_items,
            ],
            [
                "Total Amount",
                report.overview.total_amount,
            ],
            [
                "Average Receipt",
                report.overview.average_receipt,
            ],
        ]

        for row in overview_rows:
            worksheet.append(row)

        current_row = worksheet.max_row + 2

        worksheet.cell(
            current_row,
            1,
            "Currency Summary",
        ).font = header_font

        currency_header_row = current_row + 1

        currency_headers = [
            "Currency",
            "Receipt Count",
            "Total Amount",
            "Average Receipt",
            "Subtotal",
            "Tax",
            "Discount",
        ]

        for col, header in enumerate(
            currency_headers,
            start=1,
        ):
            cell = worksheet.cell(
                currency_header_row,
                col,
                header,
            )
            cell.font = header_font
            cell.fill = header_fill

        current_row = currency_header_row + 1

        for currency in report.currencies:
            worksheet.append([
                currency.currency,
                currency.receipt_count,
                currency.total_amount,
                currency.average_receipt,
                currency.subtotal,
                currency.total_tax,
                currency.total_discount,
            ])

        current_row = worksheet.max_row + 2

        worksheet.cell(
            current_row,
            1,
            "Processing Activity",
        ).font = header_font

        processing_header_row = current_row + 1

        processing_headers = [
            "Date",
            "Receipts",
            "Items",
        ]

        for col, header in enumerate(
            processing_headers,
            start=1,
        ):
            cell = worksheet.cell(
                processing_header_row,
                col,
                header,
            )
            cell.font = header_font
            cell.fill = header_fill

        for item in report.processing_trend:
            worksheet.append([
                item.date,
                item.receipts,
                item.items,
            ])

        current_row = worksheet.max_row + 2

        worksheet.cell(
            current_row,
            1,
            "Payment Methods",
        ).font = header_font

        payment_header_row = current_row + 1

        payment_headers = [
            "Method",
            "Count",
        ]

        for col, header in enumerate(
            payment_headers,
            start=1,
        ):
            cell = worksheet.cell(
                payment_header_row,
                col,
                header,
            )
            cell.font = header_font
            cell.fill = header_fill

        for item in report.payment_methods:
            worksheet.append([
                item.method,
                item.count,
            ])

        current_row = worksheet.max_row + 2

        worksheet.cell(
            current_row,
            1,
            "Merchant Categories",
        ).font = header_font

        category_header_row = current_row + 1

        category_headers = [
            "Category",
            "Count",
        ]

        for col, header in enumerate(
            category_headers,
            start=1,
        ):
            cell = worksheet.cell(
                category_header_row,
                col,
                header,
            )
            cell.font = header_font
            cell.fill = header_fill

        for item in report.merchant_categories:
            worksheet.append([
                item.category,
                item.count,
            ])

        current_row = worksheet.max_row + 2

        worksheet.cell(
            current_row,
            1,
            "Quality & QR",
        ).font = header_font

        quality_rows = [
            [
                "Successful Analyses",
                report.vision.successful,
            ],
            [
                "Failed Analyses",
                report.vision.failed,
            ],
            [
                "Success Rate",
                f"{report.vision.success_rate}%",
            ],
            [
                "Receipts With QR",
                report.qr.detected_receipts,
            ],
            [
                "Receipts Without QR",
                report.qr.without_qr,
            ],
        ]

        for row in quality_rows:
            worksheet.append(row)

        current_row = worksheet.max_row + 2

        worksheet.cell(
            current_row,
            1,
            "Receipt Details",
        ).font = header_font

        receipt_header_row = current_row + 1

        receipt_headers = [
            "ID",
            "Merchant",
            "Category",
            "Invoice",
            "Date",
            "Time",
            "Currency",
            "Subtotal",
            "Tax",
            "Discount",
            "Total",
            "Payment",
            "QR Detected",
            "QR Type",
            "QR Data",
            "Items Count",
        ]

        for col, header in enumerate(
            receipt_headers,
            start=1,
        ):
            cell = worksheet.cell(
                receipt_header_row,
                col,
                header,
            )
            cell.font = header_font
            cell.fill = header_fill

        for receipt in report.receipts:
            worksheet.append([
                receipt.id,
                receipt.merchant_name or "",
                receipt.merchant_category or "",
                receipt.invoice_number or "",
                receipt.receipt_date or "",
                receipt.receipt_time or "",
                receipt.currency or "",
                receipt.subtotal,
                receipt.tax,
                receipt.discount,
                receipt.total,
                receipt.payment_method or "",
                "Yes" if receipt.qr_detected else "No",
                receipt.qr_type or "",
                receipt.qr_data or "",
                len(receipt.items),
            ])

        widths = {
            "A": 10,
            "B": 24,
            "C": 20,
            "D": 18,
            "E": 15,
            "F": 12,
            "G": 12,
            "H": 14,
            "I": 12,
            "J": 12,
            "K": 14,
            "L": 16,
            "M": 14,
            "N": 12,
            "O": 42,
            "P": 12,
        }

        for column, width in widths.items():
            worksheet.column_dimensions[
                column
            ].width = width

        worksheet.freeze_panes = "A1"

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        filename = (
            f"smartreceiptai-report-"
            f"{report.start_date}-"
            f"{report.end_date}.xlsx"
        )

        return StreamingResponse(
            output,
            media_type=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition": (
                    f'attachment; filename="{filename}"'
                )
            },
        )

    except ValueError as exc:
        raise HTTPException(
            status_code=400,
            detail=str(exc),
        ) from exc

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail="Failed to export report to Excel.",
        ) from exc


@router.get(
    "/export/pdf",
)
def export_report_pdf(
    start_date: date,
    end_date: date,
    db: Session = Depends(get_db),
    current_user: UserModel = Depends(
        get_current_user
    ),
):
    try:
        report = get_report_data(
            db=db,
            current_user=current_user,
            start_date=start_date,
            end_date=end_date,
        )

        output = BytesIO()

        document = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=12 * mm,
            leftMargin=12 * mm,
            topMargin=12 * mm,
            bottomMargin=12 * mm,
            title="SmartReceiptAI Report",
            author="SmartReceiptAI",
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "ReportTitle",
            parent=styles["Title"],
            fontSize=18,
            leading=22,
            alignment=TA_LEFT,
            spaceAfter=8,
        )

        subtitle_style = ParagraphStyle(
            "ReportSubtitle",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#667085"),
            spaceAfter=12,
        )

        heading_style = ParagraphStyle(
            "ReportHeading",
            parent=styles["Heading2"],
            fontSize=12,
            leading=15,
            textColor=colors.HexColor("#1F2937"),
            spaceBefore=10,
            spaceAfter=7,
        )

        small_style = ParagraphStyle(
            "ReportSmall",
            parent=styles["Normal"],
            fontSize=7.5,
            leading=9,
        )

        small_bold_style = ParagraphStyle(
            "ReportSmallBold",
            parent=small_style,
            fontName="Helvetica-Bold",
        )

        story = []

        story.append(
            Paragraph(
                "SmartReceiptAI Report",
                title_style,
            )
        )

        story.append(
            Paragraph(
                f"Period: {report.start_date} → "
                f"{report.end_date}<br/>"
                f"Generated: {report.generated_at}",
                subtitle_style,
            )
        )

        story.append(
            Paragraph(
                "Overview",
                heading_style,
            )
        )

        overview_table = Table(
            [
                [
                    "Total Receipts",
                    str(
                        report.overview.total_receipts
                    ),
                    "Total Items",
                    str(
                        report.overview.total_items
                    ),
                    "Total Amount",
                    (
                        str(
                            report.overview.total_amount
                        )
                        if len(report.currencies) == 1
                        else "Multiple currencies"
                    ),
                    "Average",
                    (
                        str(
                            report.overview.average_receipt
                        )
                        if len(report.currencies) == 1
                        else "Multiple currencies"
                    ),
                ]
            ],
            colWidths=[
                28 * mm,
                18 * mm,
                24 * mm,
                18 * mm,
                24 * mm,
                28 * mm,
                20 * mm,
                28 * mm,
            ],
        )

        overview_table.setStyle(
            TableStyle([
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, -1),
                    colors.HexColor("#F5F7FA"),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, -1),
                    colors.HexColor("#1F2937"),
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.HexColor("#D9DEE7"),
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    "Helvetica",
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (0, 0),
                    "Helvetica-Bold",
                ),
                (
                    "FONTNAME",
                    (2, 0),
                    (2, 0),
                    "Helvetica-Bold",
                ),
                (
                    "FONTNAME",
                    (4, 0),
                    (4, 0),
                    "Helvetica-Bold",
                ),
                (
                    "FONTNAME",
                    (6, 0),
                    (6, 0),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ])
        )

        story.append(
            overview_table
        )

        story.append(
            Paragraph(
                "Currency Summary",
                heading_style,
            )
        )

        currency_data = [
            [
                "Currency",
                "Receipts",
                "Total",
                "Average",
                "Subtotal",
                "Tax",
                "Discount",
            ]
        ]

        for currency in report.currencies:
            currency_data.append([
                currency.currency,
                str(currency.receipt_count),
                str(currency.total_amount),
                str(currency.average_receipt),
                str(currency.subtotal),
                str(currency.total_tax),
                str(currency.total_discount),
            ])

        if len(currency_data) == 1:
            currency_data.append([
                "No data",
                "",
                "",
                "",
                "",
                "",
                "",
            ])

        currency_table = Table(
            currency_data,
            repeatRows=1,
            colWidths=[
                28 * mm,
                20 * mm,
                28 * mm,
                28 * mm,
                28 * mm,
                24 * mm,
                28 * mm,
            ],
        )

        currency_table.setStyle(
            TableStyle([
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#EDF0F4"),
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.HexColor("#D9DEE7"),
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7.5,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
            ])
        )

        story.append(currency_table)

        story.append(
            Paragraph(
                "Processing & QR Summary",
                heading_style,
            )
        )

        quality_data = [
            [
                "Successful",
                str(
                    report.vision.successful
                ),
                "Failed",
                str(
                    report.vision.failed
                ),
                "Success Rate",
                f"{report.vision.success_rate}%",
                "With QR",
                str(
                    report.qr.detected_receipts
                ),
                "Without QR",
                str(
                    report.qr.without_qr
                ),
            ]
        ]

        quality_table = Table(
            quality_data,
            colWidths=[
                24 * mm,
                18 * mm,
                18 * mm,
                18 * mm,
                24 * mm,
                22 * mm,
                18 * mm,
                18 * mm,
                24 * mm,
                22 * mm,
            ],
        )

        quality_table.setStyle(
            TableStyle([
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, -1),
                    colors.HexColor("#F5F7FA"),
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.HexColor("#D9DEE7"),
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7.5,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    "Helvetica",
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
            ])
        )

        story.append(quality_table)

        story.append(
            Paragraph(
                "Receipt Details",
                heading_style,
            )
        )

        receipt_data = [[
            "ID",
            "Merchant",
            "Invoice",
            "Date",
            "Currency",
            "Total",
            "Payment",
            "QR",
        ]]

        for receipt in report.receipts:

            merchant = (
                receipt.merchant_name
                or "Unknown"
            )

            invoice = (
                receipt.invoice_number
                or "—"
            )

            receipt_date = (
                receipt.receipt_date
                or "—"
            )

            currency = (
                receipt.currency
                or "—"
            )

            total = (
                str(receipt.total)
                if receipt.total is not None
                else "—"
            )

            payment = (
                receipt.payment_method
                or "—"
            )

            qr = (
                "Detected"
                if receipt.qr_detected
                else "No QR"
            )

            receipt_data.append([
                Paragraph(
                    str(receipt.id),
                    small_style,
                ),
                Paragraph(
                    merchant,
                    small_style,
                ),
                Paragraph(
                    invoice,
                    small_style,
                ),
                Paragraph(
                    receipt_date,
                    small_style,
                ),
                Paragraph(
                    currency,
                    small_style,
                ),
                Paragraph(
                    total,
                    small_style,
                ),
                Paragraph(
                    payment,
                    small_style,
                ),
                Paragraph(
                    qr,
                    small_style,
                ),
            ])

        if len(receipt_data) == 1:
            receipt_data.append([
                "",
                "No receipts found.",
                "",
                "",
                "",
                "",
                "",
                "",
            ])

        receipt_table = Table(
            receipt_data,
            repeatRows=1,
            colWidths=[
                12 * mm,
                55 * mm,
                32 * mm,
                28 * mm,
                22 * mm,
                25 * mm,
                28 * mm,
                20 * mm,
            ],
        )

        receipt_table.setStyle(
            TableStyle([
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#EDF0F4"),
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.HexColor("#D9DEE7"),
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7.5,
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
            ])
        )

        story.append(
            receipt_table
        )

        document.build(story)

        output.seek(0)

        filename = (
            f"smartreceiptai-report-"
            f"{report.start_date}-"
            f"{report.end_date}.pdf"
        )

        return StreamingResponse(
            output,
            media_type="application/pdf",
            headers={
                "Content-Disposition": (
                    f'attachment; filename="{filename}"'
                )
            },
        )

    except ValueError as exc:
        raise HTTPException(
            status_code=400,
            detail=str(exc),
        ) from exc

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail="Failed to export report to PDF.",
        ) from exc